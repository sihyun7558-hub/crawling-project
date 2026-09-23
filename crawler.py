import customtkinter as ctk
from playwright.sync_api import sync_playwright
import time
import threading
import os
import requests
from PIL import Image
from io import BytesIO
from datetime import datetime
import webbrowser
import csv
from tkinter import filedialog
import tkinter as tk
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import re
from collections import Counter

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

STOP_WORDS = {
    "the", "a", "an", "and", "or", "of", "in", "is", "it", "to", "for",
    "with", "on", "at", "by", "from", "as", "be", "was", "are", "were",
    "이", "그", "저", "것", "수", "더", "및", "등", "를", "을", "의", "가",
    "에", "은", "는", "이", "로", "도", "만", "한", "하", "대", "색", "형",
    "용", "남", "여", "성", "전", "후", "상", "하", "중", "무", "유", "빅",
    "new", "sale", "best", "top", "set", "no", "x",
}


class MusinsaApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("MUSINSA Ranking Crawler v5.0")
        self.geometry("1400x900")
        self.minsize(1100, 700)

        self.current_dir = os.getcwd()
        self.download_folder = os.path.join(self.current_dir, "musinsa_images")
        os.makedirs(self.download_folder, exist_ok=True)

        self.today_str = datetime.now().strftime("%Y%m%d")
        self._stop_flag = False

        # ── 헤더 ──────────────────────────────────────────────
        ctk.CTkLabel(
            self, text="무신사 랭킹 크롤러 v5.0",
            font=ctk.CTkFont(size=24, weight="bold")
        ).pack(pady=(15, 8))

        # ── 컨트롤 영역 ───────────────────────────────────────
        ctrl = ctk.CTkFrame(self, fg_color="transparent")
        ctrl.pack(fill="x", padx=24, pady=0)

        # 폴더 행
        fr = ctk.CTkFrame(ctrl)
        fr.pack(fill="x", pady=(0, 5))
        ctk.CTkLabel(fr, text="저장 폴더:", font=ctk.CTkFont(size=12, weight="bold"), width=85).pack(side="left", padx=10)
        self.folder_display = ctk.CTkLabel(fr, text=self.download_folder,
                                            font=ctk.CTkFont(size=11), text_color="#00D9FF", anchor="w")
        self.folder_display.pack(side="left", padx=8, fill="x", expand=True)
        ctk.CTkButton(fr, text="폴더 선택", command=self.select_folder,
                      width=130, height=36, fg_color="#1f538d").pack(side="right", padx=10, pady=8)

        # URL 행
        ur = ctk.CTkFrame(ctrl)
        ur.pack(fill="x", pady=(0, 5))
        ctk.CTkLabel(ur, text="링크 입력:", font=ctk.CTkFont(size=12, weight="bold"), width=85).pack(side="left", padx=10)

        self.url_entry = ctk.CTkEntry(ur, placeholder_text="무신사 랭킹 주소를 붙여넣으세요", height=38)
        self.url_entry.pack(side="left", padx=8, pady=8, fill="x", expand=True)

        # Ctrl+V 수동 바인딩 (return "break" 으로 기본 동작 차단 후 직접 처리)
        self.url_entry.bind("<Control-v>", self._on_ctrl_v)
        self.url_entry.bind("<Control-V>", self._on_ctrl_v)
        self.url_entry.bind("<Control-a>", self._on_ctrl_a)
        self.url_entry.bind("<Control-A>", self._on_ctrl_a)
        self.url_entry.bind("<Button-3>",  self._show_ctx_menu)

        self.stop_button = ctk.CTkButton(
            ur, text="⏹  정지", command=self.stop_crawling,
            width=110, height=38, fg_color="#7a1e1e", state="disabled"
        )
        self.stop_button.pack(side="right", padx=4, pady=8)

        self.run_button = ctk.CTkButton(
            ur, text="▶  추출 시작", command=self.start_crawling,
            width=130, height=38, fg_color="#1a6b3a"
        )
        self.run_button.pack(side="right", padx=4, pady=8)

        # 순위 행
        rr = ctk.CTkFrame(ctrl)
        rr.pack(fill="x", pady=(0, 5))
        ctk.CTkLabel(rr, text="순위 범위:", font=ctk.CTkFont(size=12, weight="bold"), width=85).pack(side="left", padx=10)
        ctk.CTkLabel(rr, text="시작", font=ctk.CTkFont(size=11)).pack(side="left", padx=(8, 2))
        self.rank_start_entry = ctk.CTkEntry(rr, width=72, height=34)
        self.rank_start_entry.insert(0, "1")
        self.rank_start_entry.pack(side="left", padx=(0, 6), pady=8)
        ctk.CTkLabel(rr, text="~", font=ctk.CTkFont(size=14, weight="bold")).pack(side="left")
        ctk.CTkLabel(rr, text="끝", font=ctk.CTkFont(size=11)).pack(side="left", padx=(6, 2))
        self.rank_end_entry = ctk.CTkEntry(rr, width=72, height=34)
        self.rank_end_entry.insert(0, "50")
        self.rank_end_entry.pack(side="left", padx=(0, 14), pady=8)
        for lbl, s, e in [("TOP 10", 1, 10), ("TOP 20", 1, 20), ("TOP 50", 1, 50), ("TOP 100", 1, 100)]:
            ctk.CTkButton(rr, text=lbl, command=lambda a=s, b=e: self.set_rank_range(a, b),
                          width=76, height=34, fg_color="#2b5278").pack(side="left", padx=3, pady=8)
        ctk.CTkLabel(rr, text="※ 100위 초과 시 시간이 길어질 수 있습니다",
                     font=ctk.CTkFont(size=10), text_color="#777777").pack(side="left", padx=14)

        # 진행 바 행 (fill="x" 로 위 칸들과 동일 폭)
        pr = ctk.CTkFrame(ctrl, fg_color="transparent")
        pr.pack(fill="x", pady=(4, 10))
        self.progress_label = ctk.CTkLabel(pr, text="대기 중... (0%)", font=ctk.CTkFont(size=11))
        self.progress_label.pack(anchor="w", padx=2)
        self.progress_bar = ctk.CTkProgressBar(pr)
        self.progress_bar.set(0)
        self.progress_bar.pack(fill="x", padx=2, pady=(2, 0))

        # ── 메인 영역 (grid 2:3 비율) ─────────────────────────
        main = ctk.CTkFrame(self, fg_color="transparent")
        main.pack(fill="both", expand=True, padx=24, pady=(0, 12))
        main.columnconfigure(0, weight=2)
        main.columnconfigure(1, weight=3)
        main.rowconfigure(0, weight=1)

        # 왼쪽 패널 (로그 + 키워드)
        left = ctk.CTkFrame(main)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        left.rowconfigure(0, weight=3)
        left.rowconfigure(1, weight=1)
        left.columnconfigure(0, weight=1)

        self.output_text = ctk.CTkTextbox(left, font=("Pretendard", 12))
        self.output_text.grid(row=0, column=0, sticky="nsew", padx=5, pady=(5, 4))

        # 키워드 패널
        kw_outer = ctk.CTkFrame(left, fg_color="#1a1a2e", corner_radius=10)
        kw_outer.grid(row=1, column=0, sticky="nsew", padx=5, pady=(0, 5))
        kw_outer.rowconfigure(1, weight=1)
        kw_outer.columnconfigure(0, weight=1)

        kw_hdr = ctk.CTkFrame(kw_outer, fg_color="transparent")
        kw_hdr.grid(row=0, column=0, sticky="ew", padx=10, pady=(8, 2))
        ctk.CTkLabel(kw_hdr, text="🔥 키워드 분석 TOP 10",
                     font=ctk.CTkFont(size=13, weight="bold"), text_color="#FFD700").pack(side="left")
        self.keyword_update_label = ctk.CTkLabel(kw_hdr, text="",
                                                  font=ctk.CTkFont(size=10), text_color="#666666")
        self.keyword_update_label.pack(side="right")

        self.keyword_scroll = ctk.CTkScrollableFrame(kw_outer, fg_color="transparent", height=155)
        self.keyword_scroll.grid(row=1, column=0, sticky="nsew", padx=8, pady=(0, 8))
        self.keyword_content = self.keyword_scroll

        ctk.CTkLabel(self.keyword_content, text="크롤링 완료 후 키워드가 표시됩니다",
                     font=ctk.CTkFont(size=11), text_color="#555555").pack(pady=8)

        # 오른쪽 패널 (갤러리)
        right = ctk.CTkFrame(main)
        right.grid(row=0, column=1, sticky="nsew")
        ctk.CTkLabel(right, text="GALLERY", font=ctk.CTkFont(size=16, weight="bold")).pack(pady=10)
        self.scroll_frame = ctk.CTkScrollableFrame(right, fg_color="#1e1e1e")
        self.scroll_frame.pack(fill="both", expand=True, padx=5, pady=5)

        self.image_objects = []
        self.gallery_columns = 3
        self.products_data = []

        # 우클릭 메뉴
        self._ctx_menu = tk.Menu(self, tearoff=0)
        self._ctx_menu.add_command(label="붙여넣기 (Ctrl+V)", command=self._paste_url)
        self._ctx_menu.add_command(label="복사 (Ctrl+C)",     command=self._copy_url)
        self._ctx_menu.add_command(label="전체 선택 (Ctrl+A)", command=self._select_all_url)
        self._ctx_menu.add_separator()
        self._ctx_menu.add_command(label="지우기", command=lambda: self.url_entry.delete(0, "end"))

    # ── 붙여넣기 핸들러 ───────────────────────────────────────
    def _on_ctrl_v(self, event=None):
        try:
            clip = self.clipboard_get()
            try:
                s = self.url_entry.index(tk.SEL_FIRST)
                e = self.url_entry.index(tk.SEL_LAST)
                self.url_entry.delete(s, e)
                self.url_entry.insert(s, clip.strip())
            except tk.TclError:
                pos = self.url_entry.index(tk.INSERT)
                self.url_entry.insert(pos, clip.strip())
        except Exception:
            pass
        return "break"

    def _on_ctrl_a(self, event=None):
        try:
            self.url_entry.select_range(0, "end")
        except Exception:
            pass
        return "break"

    def _show_ctx_menu(self, event):
        self._ctx_menu.tk_popup(event.x_root, event.y_root)

    def _paste_url(self):       self._on_ctrl_v()
    def _copy_url(self):
        try:
            self.clipboard_clear(); self.clipboard_append(self.url_entry.get())
        except Exception: pass
    def _select_all_url(self):  self._on_ctrl_a()

    # ── 순위 빠른 설정 ────────────────────────────────────────
    def set_rank_range(self, s, e):
        self.rank_start_entry.delete(0, "end"); self.rank_start_entry.insert(0, str(s))
        self.rank_end_entry.delete(0, "end");   self.rank_end_entry.insert(0, str(e))

    # ── 폴더 선택 ─────────────────────────────────────────────
    def select_folder(self):
        folder = filedialog.askdirectory(title="저장 폴더 선택", initialdir=self.current_dir)
        if folder:
            self.download_folder = folder
            self.folder_display.configure(text=folder)
            self.add_result(f"✅ 저장 폴더: {folder}")
            os.makedirs(folder, exist_ok=True)

    # ── 이미지 팝업 ───────────────────────────────────────────
    def open_image_popup(self, img_path, brand, title, price):
        popup = ctk.CTkToplevel(self)
        popup.title(f"{brand} - {title}")
        popup.geometry("500x700")
        img = Image.open(img_path); img.thumbnail((480, 600))
        ctk_img = ctk.CTkImage(light_image=img, dark_image=img, size=img.size)
        ctk.CTkLabel(popup, image=ctk_img).pack(pady=10)
        ctk.CTkLabel(popup, text=f"{brand}\n{title}\n{price}원",
                     font=ctk.CTkFont(size=12), justify="center").pack(pady=5)
        ctk.CTkButton(popup, text="상품 페이지 열기",
                      command=lambda: webbrowser.open(f"https://www.musinsa.com/app/goods/{title}")).pack(pady=10)

    # ── 갤러리 추가 ───────────────────────────────────────────
    def add_image_to_gallery(self, img_url, rank_num, brand, title, price, img_path):
        try:
            r = requests.get(img_url, timeout=5)
            img_data = Image.open(BytesIO(r.content)); img_data.thumbnail((150, 200))
            ctk_img = ctk.CTkImage(light_image=img_data, dark_image=img_data, size=img_data.size)
            idx = len(self.image_objects)
            row = (idx // self.gallery_columns) * 2
            col = idx % self.gallery_columns
            lbl = ctk.CTkLabel(self.scroll_frame, image=ctk_img, text="")
            lbl.grid(row=row, column=col, padx=5, pady=5)
            lbl.bind("<Button-1>", lambda e, p=img_path, b=brand, t=title, pr=price:
                     self.open_image_popup(p, b, t, pr))
            ctk.CTkLabel(self.scroll_frame,
                         text=f"{rank_num}위\n{brand}\n{title}\n{price}원",
                         font=ctk.CTkFont(size=10), wraplength=150, justify="center"
                         ).grid(row=row+1, column=col, padx=5, pady=(0, 10))
            self.image_objects.append(ctk_img)
        except Exception as e:
            print(f"갤러리 추가 실패: {e}")

    # ── 키워드 분석 ───────────────────────────────────────────
    def analyze_and_show_keywords(self, titles):
        tokens = []
        for t in titles:
            tokens.extend([w.lower() for w in re.findall(r'[A-Za-z]{2,}', t)])
            tokens.extend(re.findall(r'[가-힣]{2,}', t))
        filtered = [t for t in tokens if t not in STOP_WORDS]
        top10 = Counter(filtered).most_common(10)
        self.after(0, lambda: self._render_keywords(top10, len(titles)))

    def _render_keywords(self, top10, total):
        for w in self.keyword_content.winfo_children():
            w.destroy()
        if not top10:
            ctk.CTkLabel(self.keyword_content, text="키워드 없음",
                         font=ctk.CTkFont(size=11), text_color="#555555").pack(pady=8)
            return
        self.keyword_update_label.configure(text=f"상품 {total}개 분석")
        max_c = top10[0][1]
        colors = ["#FF4444","#FF6B35","#FF9500","#FFB800","#FFD700",
                  "#A8E063","#56CCF2","#6C9BEF","#9B59B6","#808080"]
        for i, (kw, cnt) in enumerate(top10):
            row = ctk.CTkFrame(self.keyword_content, fg_color="transparent")
            row.pack(fill="x", pady=2)
            c = colors[i] if i < len(colors) else "#808080"
            ctk.CTkLabel(row, text=f"{i+1:2d}", font=ctk.CTkFont(size=11, weight="bold"),
                         text_color=c, width=24).pack(side="left", padx=(0, 5))
            ctk.CTkLabel(row, text=kw, font=ctk.CTkFont(size=12, weight="bold"),
                         text_color="#EEEEEE", width=95, anchor="w").pack(side="left", padx=(0, 6))
            bf = ctk.CTkFrame(row, fg_color="#333333", width=148, height=13, corner_radius=4)
            bf.pack(side="left", padx=(0, 6))
            bf.pack_propagate(False)
            ctk.CTkFrame(bf, fg_color=c, width=max(4, int(148*cnt/max_c)),
                         height=13, corner_radius=4).place(x=0, y=0)
            ctk.CTkLabel(row, text=f"{cnt}회", font=ctk.CTkFont(size=11),
                         text_color="#AAAAAA", width=38).pack(side="left")

    # ── 로그 ──────────────────────────────────────────────────
    def add_result(self, text):
        self.after(0, lambda: (
            self.output_text.insert("end", text + "\n"),
            self.output_text.see("end")
        ))

    # ── 진행 바 ───────────────────────────────────────────────
    def update_progress_smoothly(self, target, duration=0.3):
        cur = self.progress_bar.get()
        steps = 15
        inc = (target - cur) / steps
        for _ in range(steps):
            cur += inc
            v = max(0, min(cur, 1.0))
            self.progress_bar.set(v)
            self.after(0, lambda x=v: self.progress_label.configure(text=f"수집 중... ({int(x*100)}%)"))
            time.sleep(duration / steps)

    # ── 이미지 저장 ───────────────────────────────────────────
    def download_image(self, img_url, brand, title, price, rank_num):
        try:
            if not img_url or "data:" in img_url: return None
            fn = f"{self.today_str}_{rank_num}위_{brand.replace('/','_').replace(' ','_')}_" \
                 f"{title.replace('/','_').replace(' ','_')}_{price.replace(',','')}.jpg"
            path = os.path.join(self.download_folder, fn)
            r = requests.get(img_url, stream=True, timeout=10)
            if r.status_code == 200:
                with open(path, 'wb') as f:
                    for chunk in r.iter_content(1024): f.write(chunk)
                return path
        except Exception: pass
        return None

    # ── 엑셀 저장 ─────────────────────────────────────────────
    def save_to_excel(self, excel_file):
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "무신사 랭킹"
            hf = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            for ci, h in enumerate(["순위","브랜드","상품명","가격","파일명"], 1):
                c = ws.cell(row=1, column=ci, value=h)
                c.fill = hf
                c.font = Font(bold=True, color="FFFFFF", size=12)
                c.alignment = Alignment(horizontal="center", vertical="center")
            for ri, row in enumerate(self.products_data, 2):
                for ci, val in enumerate(row, 1):
                    ws.cell(row=ri, column=ci, value=val).alignment = Alignment(horizontal="center", vertical="center")
            for col, w in zip("ABCDE", [10,20,40,15,50]):
                ws.column_dimensions[col].width = w
            wb.save(excel_file)
            self.add_result(f"📊 엑셀 저장: {os.path.basename(excel_file)}")
            return True
        except Exception as e:
            self.add_result(f"❌ 엑셀 오류: {e}"); return False

    # ── 정지 ──────────────────────────────────────────────────
    def stop_crawling(self):
        self._stop_flag = True
        self.add_result("⏹ 정지 요청 — 현재 항목 처리 후 중단됩니다...")
        self.stop_button.configure(state="disabled")

    # ── 크롤링 시작 ───────────────────────────────────────────
    def start_crawling(self):
        target_url = self.url_entry.get().strip()
        if not target_url.startswith("http"):
            self.add_result("⚠️ 올바른 URL을 입력해주세요."); return
        try:
            rank_start = int(self.rank_start_entry.get().strip())
            rank_end   = int(self.rank_end_entry.get().strip())
            if rank_start < 1 or rank_end < rank_start: raise ValueError
        except ValueError:
            self.add_result("⚠️ 순위 범위를 올바르게 입력해주세요."); return

        self.output_text.delete("1.0", "end")
        for w in self.scroll_frame.winfo_children(): w.destroy()
        for w in self.keyword_content.winfo_children(): w.destroy()
        ctk.CTkLabel(self.keyword_content, text="크롤링 완료 후 키워드가 표시됩니다",
                     font=ctk.CTkFont(size=11), text_color="#555555").pack(pady=8)
        self.keyword_update_label.configure(text="")
        self.image_objects = []; self.products_data = []; self._stop_flag = False

        self.run_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.progress_bar.set(0)
        self.progress_label.configure(text="시작 중... (0%)")

        threading.Thread(target=self.run_scraper, args=(target_url, rank_start, rank_end), daemon=True).start()

    # ── 크롤러 핵심 ───────────────────────────────────────────
    def run_scraper(self, target_url, rank_start, rank_end):
        stopped_early = False
        try:
            csv_file   = os.path.join(self.download_folder, f"{self.today_str}_musinsa.csv")
            excel_file = os.path.join(self.download_folder, f"{self.today_str}_musinsa.xlsx")
            self.add_result(f"🎯 수집 범위: {rank_start}위 ~ {rank_end}위 ({rank_end-rank_start+1}개)")

            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                page = browser.new_page()
                page.goto(url=target_url, wait_until="networkidle")

                collected = {}
                target_count = rank_end - rank_start + 1
                max_scrolls  = max(30, rank_end // 3)

                for _ in range(max_scrolls):
                    if self._stop_flag: stopped_early = True; break
                    self.update_progress_smoothly(min(len(collected) / target_count, 0.95))
                    if len(collected) >= target_count: break

                    for item in page.locator('div.gtm-view-item-list').all():
                        if self._stop_flag: stopped_early = True; break
                        try:
                            rt = item.locator('span.text-etc_11px_semibold.text-black.font-pretendard').first.inner_text()
                            if rt.isdigit():
                                rn = int(rt)
                                if rank_start <= rn <= rank_end and rn not in collected:
                                    brand = item.get_attribute("data-item-brand") or "브랜드없음"
                                    title = item.locator('p.text-body_13px_reg').first.inner_text().strip()
                                    price = item.get_attribute("data-price") or "0"
                                    img_el = item.locator('img').first
                                    img_url = img_el.get_attribute("src") or img_el.get_attribute("data-original")
                                    if not img_url: continue

                                    img_path = self.download_image(img_url, brand, title, price, rn)
                                    if img_path:
                                        self.after(0, self.add_image_to_gallery,
                                                   img_url, rn, brand, title, price, img_path)

                                    info = [rn, brand, title, price, os.path.basename(img_path) if img_path else ""]
                                    self.products_data.append(info)
                                    self.add_result(f"[{rn}위] {brand} - {title} ({price}원) | {'📸 저장완료' if img_path else '❌ 저장실패'}")
                                    collected[rn] = info
                                    if len(collected) >= target_count: break
                        except Exception: continue

                    if stopped_early: break
                    page.mouse.wheel(0, 3500)
                    time.sleep(1.2)

                browser.close()

            # 수집된 데이터 저장 (정지해도 저장)
            if self.products_data:
                with open(csv_file, 'w', newline='', encoding='utf-8-sig') as f:
                    csv.writer(f).writerows([["순위","브랜드","상품명","가격","파일명"]] + self.products_data)
                self.save_to_excel(excel_file)
                titles = [p[2] for p in self.products_data if len(p) > 2]
                if titles:
                    self.add_result(f"\n🔍 키워드 분석 중... ({len(titles)}개)")
                    self.analyze_and_show_keywords(titles)

            n = len(self.products_data)
            if stopped_early:
                self.after(0, lambda: (
                    self.progress_label.configure(text=f"⏹ 중단됨 ({n}개 수집)"),
                    self.progress_bar.set(n / max(rank_end-rank_start+1, 1))
                ))
                self.add_result(f"\n⏹ 중단! {n}개 수집 후 저장 완료.")
            else:
                self.after(0, lambda: (
                    self.progress_label.configure(text="완료! (100%)"),
                    self.progress_bar.set(1.0)
                ))
                self.add_result(f"\n✅ 완료! 총 {n}개 ({rank_start}위~{rank_end}위)")
            self.add_result(f"📁 {self.download_folder}")

        except Exception as e:
            self.add_result(f"\n❌ 오류: {e}")
        finally:
            self.after(0, lambda: (
                self.run_button.configure(state="normal"),
                self.stop_button.configure(state="disabled")
            ))


if __name__ == "__main__":
    app = MusinsaApp()
    app.mainloop()
